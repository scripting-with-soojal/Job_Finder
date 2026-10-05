import csv
import os
import tempfile
import unittest

from openpyxl import load_workbook

from job_finder.cli import load_boards, main
from job_finder.pipeline import run_pipeline
from job_finder.sources import adzuna, greenhouse


class FakeResponse:
    def __init__(self, payload, status=200):
        self._payload, self.status_code = payload, status

    def json(self):
        return self._payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class FakeSession:
    """Routes URLs to canned responses; records calls."""

    def __init__(self, routes):
        self.routes, self.calls = routes, []

    def get(self, url, params=None, timeout=None):
        self.calls.append((url, params))
        for fragment, resp in self.routes.items():
            if fragment in url:
                return resp() if callable(resp) else resp
        raise AssertionError(f"unexpected URL {url}")


ADZUNA_PAGE = {"results": [{
    "company": {"display_name": "Acme"}, "title": "Data Engineer",
    "location": {"display_name": "Mumbai"},
    "description": "3-5 years. Python, SQL, Airflow",
    "redirect_url": "https://adz/1", "salary_min": 100, "salary_max": 200,
    "created": "2026-10-01",
}]}
GREENHOUSE = {"jobs": [
    {"title": "Senior Data Engineer", "location": {"name": "London"},
     "content": "&lt;p&gt;5+ years of &lt;b&gt;Snowflake&lt;/b&gt;&lt;/p&gt;",
     "absolute_url": "https://gh/1", "updated_at": "2026-10-02"},
    {"title": "Product Designer", "location": {"name": "London"},
     "content": "", "absolute_url": "https://gh/2", "updated_at": "x"},
]}


class TestSources(unittest.TestCase):
    def test_adzuna_stops_on_empty_page(self):
        pages = iter([FakeResponse(ADZUNA_PAGE), FakeResponse({"results": []})])
        s = FakeSession({"adzuna": lambda: next(pages)})
        jobs = adzuna.fetch_jobs(s, "data engineer", "in", "id", "key", max_pages=10, delay=0)
        self.assertEqual(len(jobs), 1)
        self.assertEqual(len(s.calls), 2)  # stopped at the empty page, not page 10
        self.assertEqual(jobs[0].salary_min, 100)

    def test_greenhouse_filters_title_and_cleans_html(self):
        s = FakeSession({"greenhouse": FakeResponse(GREENHOUSE)})
        jobs = greenhouse.fetch_jobs(s, "Stripe", "stripe", "data engineer")
        self.assertEqual([j.title for j in jobs], ["Senior Data Engineer"])
        self.assertEqual(jobs[0].description, "5+ years of Snowflake")


class TestPipeline(unittest.TestCase):
    boards = [{"company": "Stripe", "type": "greenhouse", "identifier": "stripe", "active": True}]

    def test_one_failing_source_does_not_stop_run(self):
        s = FakeSession({"adzuna": FakeResponse({}, status=500),
                         "greenhouse": FakeResponse(GREENHOUSE)})
        jobs, logs = run_pipeline("data engineer", "in", self.boards, "id", "key", session=s)
        self.assertEqual(len(jobs), 1)
        status = {l.source: l.status for l in logs}
        self.assertEqual(status, {"Adzuna": "FAILED", "Stripe": "SUCCESS"})

    def test_missing_credentials_logged_not_raised(self):
        s = FakeSession({"greenhouse": FakeResponse(GREENHOUSE)})
        _, logs = run_pipeline("data engineer", "in", self.boards, session=s)
        self.assertEqual(logs[0].status, "FAILED")
        self.assertIn("not set", logs[0].error)

    def test_inactive_board_skipped(self):
        s = FakeSession({"adzuna": FakeResponse({"results": []})})
        boards = [dict(self.boards[0], active=False)]
        _, logs = run_pipeline("x", "in", boards, "id", "key", session=s)
        self.assertEqual([l.source for l in logs], ["Adzuna"])


class TestCli(unittest.TestCase):
    def test_load_boards(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "b.csv")
            with open(p, "w") as f:
                f.write("company,type,identifier,active\nStripe,Greenhouse,stripe,TRUE\nX,greenhouse,x,false\n")
            b = load_boards(p)
        self.assertEqual(b[0], {"company": "Stripe", "type": "greenhouse", "identifier": "stripe", "active": True})
        self.assertFalse(b[1]["active"])

    def test_end_to_end_csv_output(self):
        import job_finder.pipeline as pl
        s = FakeSession({"adzuna": FakeResponse(ADZUNA_PAGE), "greenhouse": FakeResponse(GREENHOUSE)})
        orig = pl.build_session
        pl.build_session = lambda *a, **k: s
        orig_delay = adzuna.time.sleep
        adzuna.time.sleep = lambda *_: None
        os.environ["ADZUNA_APP_ID"], os.environ["ADZUNA_APP_KEY"] = "id", "key"
        try:
            with tempfile.TemporaryDirectory() as d:
                boards = os.path.join(d, "b.csv")
                with open(boards, "w") as f:
                    f.write("company,type,identifier,active\nStripe,greenhouse,stripe,true\n")
                out, log = os.path.join(d, "o", "jobs.csv"), os.path.join(d, "o", "logs.csv")
                main(["--role", "data engineer", "--country", "in", "--boards", boards,
                      "--max-pages", "1", "--out", out, "--log", log])
                rows = list(csv.DictReader(open(out)))
                logs = list(csv.DictReader(open(log)))
        finally:
            pl.build_session = orig
            adzuna.time.sleep = orig_delay
        self.assertEqual(len(rows), 2)
        by_src = {r["Source"]: r for r in rows}
        self.assertEqual(by_src["Adzuna"]["Experience Required"], "3-5 Years")
        self.assertEqual(by_src["Adzuna"]["Keywords"], "Python, SQL, Airflow")
        self.assertEqual(by_src["Greenhouse"]["Keywords"], "Snowflake")
        self.assertEqual(len(logs), 2)

    def test_end_to_end_xlsx_output(self):
        import job_finder.pipeline as pl
        s = FakeSession({"adzuna": FakeResponse(ADZUNA_PAGE), "greenhouse": FakeResponse(GREENHOUSE)})
        orig = pl.build_session
        pl.build_session = lambda *a, **k: s
        orig_delay = adzuna.time.sleep
        adzuna.time.sleep = lambda *_: None
        os.environ["ADZUNA_APP_ID"], os.environ["ADZUNA_APP_KEY"] = "id", "key"
        try:
            with tempfile.TemporaryDirectory() as d:
                boards = os.path.join(d, "b.csv")
                with open(boards, "w") as f:
                    f.write("company,type,identifier,active\nStripe,greenhouse,stripe,true\n")
                out, log = os.path.join(d, "o", "jobs.xlsx"), os.path.join(d, "o", "logs.csv")
                main(["--role", "data engineer", "--country", "in", "--boards", boards,
                      "--max-pages", "1", "--out", out, "--log", log])
                workbook = load_workbook(out, read_only=True)
                try:
                    rows = list(workbook["Jobs"].iter_rows(values_only=True))
                finally:
                    workbook.close()
        finally:
            pl.build_session = orig
            adzuna.time.sleep = orig_delay
        self.assertEqual(rows[0][0], "Source")
        self.assertEqual(len(rows), 3)
        by_src = {row[0]: row for row in rows[1:]}
        self.assertEqual(by_src["Adzuna"][4], "3-5 Years")
        self.assertEqual(by_src["Adzuna"][9], "Python, SQL, Airflow")
        self.assertEqual(by_src["Greenhouse"][9], "Snowflake")


if __name__ == "__main__":
    unittest.main()
